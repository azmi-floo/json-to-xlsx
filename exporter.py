import json
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
import gc
from itertools import islice


def batch_rows(data, batch_size=1000):
    """Generator function to yield batches of rows"""
    it = iter(data)
    while True:
        batch = list(islice(it, batch_size))
        if not batch:
            break
        yield batch


def set_field_format(worksheet, start_row, end_row, field_keys, settings):
    """Format a batch of rows to conserve memory"""
    currency_fields = settings.get("currency_fields", [])
    currency_symbol = settings.get("currency_symbol", "Rp")
    text_center_fields = settings.get("text_center_fields", [])
    text_middle_fields = settings.get("text_middle_fields", [])

    border = Border(left=Side(style="thin", color="808080"), right=Side(style="thin", color="808080"), top=Side(style="thin", color="808080"), bottom=Side(style="thin", color="808080"))

    # Process the batch of rows
    for row in worksheet.iter_rows(min_row=start_row, max_row=end_row):
        for cell, field_key in zip(row, field_keys):
            cell.border = border

            if field_key in currency_fields:
                try:
                    if cell.value is not None:
                        cell.value = float(cell.value)
                        cell.number_format = f'_({currency_symbol}* #,##0.00_);_({currency_symbol}* \(#,##0.00\);_({currency_symbol}* "-"??_);_(@_)'
                except (ValueError, TypeError):
                    pass

            if field_key in text_center_fields or field_key in text_middle_fields:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if field_key in text_center_fields:
                    cell.alignment = Alignment(horizontal="center")
                if field_key in text_middle_fields:
                    cell.alignment = Alignment(vertical="center")
                if field_key in text_center_fields and field_key in text_middle_fields:
                    cell.alignment = Alignment(horizontal="center", vertical="center")


def create_excel(json_data, output_file):
    wb = Workbook()
    wb.remove(wb.active)

    settings = json_data.get("settings", {})

    # Process each sheet
    for sheet_data in json_data["data"]:
        print(f"Processing sheet: {sheet_data['sheet']}")
        ws = wb.create_sheet(sheet_data["sheet"])

        field_keys = list(sheet_data["fields"].keys())
        for col, field_key in enumerate(field_keys, 1):
            cell = ws.cell(row=1, column=col, value=sheet_data["fields"][field_key])

            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(left=Side(style="thin", color="808080"), right=Side(style="thin", color="808080"), top=Side(style="thin", color="808080"), bottom=Side(style="thin", color="808080"))
            cell.fill = PatternFill(start_color="E6E6E6", end_color="E6E6E6", fill_type="solid")

        # Process data in batches
        batch_size = 1000
        total_rows = len(sheet_data["lists"])

        for batch_index, batch in enumerate(batch_rows(sheet_data["lists"], batch_size)):
            start_row = batch_index * batch_size + 2  # 1 is header, so +2 to skip the header parts

            # Writing the data to the output file
            for row_index, item in enumerate(batch, start_row):
                for col, field_key in enumerate(field_keys, 1):
                    ws.cell(row=row_index, column=col, value=item.get(field_key, ""))

            # Set field format
            end_row = min(start_row + len(batch) - 1, total_rows + 1)
            set_field_format(ws, start_row, end_row, field_keys, settings)

            gc.collect()
            print(f"Processed {min(end_row - 1, total_rows)} of {total_rows} rows")

        # Adjust column width
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)

            for cell in list(column)[:1000]:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass

            adjusted_width = min(max(max_length + 2, 15), 30)
            ws.column_dimensions[column_letter].width = adjusted_width

    print("Saving workbook...")
    wb.save(output_file)
    print("Done!")


def main():
    if len(sys.argv) != 3:
        print("Usage: python3 exporter.py <json_path> <output_path>")
        sys.exit(1)

    input_file = sys.argv[1]
    output_file = sys.argv[2]

    try:
        with open(input_file, "r") as f:
            json_data = json.load(f)

        create_excel(json_data, output_file)
        print(f"Excel file successfully created: {output_file}")

    except Exception as e:
        print(f"Error: {str(e)}")
        sys.exit(1)


if __name__ == "__main__":
    main()
